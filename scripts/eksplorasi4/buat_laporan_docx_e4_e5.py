#!/usr/bin/env python3
"""
buat_laporan_docx_e4_e5.py -- laporan evaluasi EKSPLORASI 4 & 5 (TrOCR) dlm
format .docx, meniru gaya "Laporan/Evaluasi Visual Eksplorasi.docx"
(landscape, ringkasan, temuan, rekomendasi, lalu per foto: overlay, tabel
per blok Kunci/E1/E2/E4/E5 berwarna, galeri kotak yg salah).

Butuh hasil bandingkan_e4_e5.py (folder --hasil: ringkasan.json, crops/).
Overlay geometri dibuat ulang di sini (tanpa model).

Contoh (dari folder scripts/eksplorasi4, python venv):
    .venv/Scripts/python buat_laporan_docx_e4_e5.py \
        --out "../../Laporan/Evaluasi Visual Eksplorasi 4 vs 5.docx"
"""

import argparse
import io
import json
import sys
from datetime import date
from pathlib import Path

import cv2
import numpy as np
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from refine_grid_eksplorasi import detect_grid_fast, draw_refined_overlay  # noqa: E402

INK = RGBColor(0x1D, 0x24, 0x33)
MUTED = RGBColor(0x5B, 0x64, 0x74)
RED = RGBColor(0xA3, 0x27, 0x1B)
WARN = RGBColor(0x8A, 0x5A, 0x00)
LINE, HEAD, OK, BAD = "D9DDE4", "EEF1F5", "E5F4EA", "FBE3E1"

JAM = {
    "1": ["07.00-08.00", "08.00-09.00", "09.00-10.00", "10.00-11.00",
          "11.00-12.00", "12.00-13.00", "13.00-14.00", "14.00-15.00"],
    "2": ["15.00-16.00", "16.00-17.00", "17.00-18.00", "18.00-19.00",
          "19.00-20.00", "20.00-21.00", "21.00-22.00", "22.00-23.00"],
    "3": ["23.00-24.00", "24.00-01.00", "01.00-02.00", "02.00-03.00",
          "03.00-04.00", "04.00-05.00", "05.00-06.00", "06.00-07.00"],
}
ROWS = [("E1", "e1"), ("E2", "e2"), ("E4", "e4_kandidat"), ("E5", "e5_kandidat")]
NAMA = {
    "e1": "E1 per-kotak · Ollama qwen2.5vl:3b",
    "e2": "E2 per-blok · Ollama qwen2.5vl:3b",
    "e4_kandidat": "E4 per-kotak · TrOCR kandidat (huruf a–k)",
    "e5_kandidat": "E5 per-blok · TrOCR kandidat (huruf a–k)",
    "e4_kandidat_az": "E4 per-kotak · TrOCR kandidat (huruf a–z, uji pertama)",
    "e5_kandidat_az": "E5 per-blok · TrOCR kandidat (huruf a–z, uji pertama)",
    "e4_bebas": "E4 per-kotak · TrOCR bebas",
    "e5_bebas": "E5 per-blok · TrOCR bebas",
}
TUNING_FOTO = {"IMG_20260901_134505", "IMG_20260921_004756"}


# ------------------------------------------------------------------ helper docx
def _set(el, tag, **attrs):
    child = el.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        el.append(child)
    for k, v in attrs.items():
        child.set(qn(k), str(v))
    return child


def style_cell(cell, fill=None, border=LINE, margin=(40, 60), valign="center"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = _set(tcPr, "w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        _set(borders, f"w:{side}", **{"w:val": "single", "w:sz": 4, "w:color": border})
    if fill:
        _set(tcPr, "w:shd", **{"w:val": "clear", "w:color": "auto", "w:fill": fill})
    mar = _set(tcPr, "w:tcMar")
    for side, v in (("top", margin[0]), ("bottom", margin[0]), ("left", margin[1]), ("right", margin[1])):
        _set(mar, f"w:{side}", **{"w:w": v, "w:type": "dxa"})
    _set(tcPr, "w:vAlign", **{"w:val": valign})


def cell_text(cell, text, bold=False, size=9, color=INK, align="left", mono=False, first=True):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER}[align]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = color
    if mono:
        r.font.name = "Consolas"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    return p


def img_stream(img_bgr, quality=85):
    ok, buf = cv2.imencode(".jpg", img_bgr, [cv2.IMWRITE_JPEG_QUALITY, quality])
    return io.BytesIO(buf.tobytes())


def bullet(doc, bold, text):
    p = doc.add_paragraph(style="List Paragraph")
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.space_after = Pt(3)
    r0 = p.add_run("•  ")
    if bold:
        r = p.add_run(bold + " ")
        r.bold = True
    p.add_run(text)
    r0.font.color.rgb = MUTED
    return p


def simple_table(doc, header, rows, widths_cm, highlight_col=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        style_cell(c, fill=HEAD)
        cell_text(c, h, bold=True, size=9, align="center" if j else "left")
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]
            fill = None
            if isinstance(v, (tuple, list)):
                v, fill = v
            style_cell(c, fill=fill)
            cell_text(c, str(v), size=9, align="center" if j else "left")
    for j, w in enumerate(widths_cm):
        for r in t.rows:
            r.cells[j].width = Cm(w)
    return t


# ------------------------------------------------------------------ data
def frac(s):
    if not s:
        return None
    a, b = s.split("/")
    return int(a), int(b)


def sheet_of(key, stem):
    return {"lembar_padat": "padat", "lembar_jarang": "jarang"}.get(key["foto"].get(stem), "?")


def norm(v):
    return v.lower() if isinstance(v, str) else v


def block_vals(block_rows, stem, shift, tag):
    out = []
    for r in block_rows:
        if r["foto"] == stem and r["shift"] == shift:
            out.append(r.get(tag) or [None] * 6)
    return out


def truth_of(block_rows, stem, shift):
    return [r["kunci"] for r in block_rows if r["foto"] == stem and r["shift"] == shift]


# ------------------------------------------------------------------ bagian dokumen
def add_title(doc, subtitle):
    doc.add_paragraph("Evaluasi Visual Eksplorasi 4 vs Eksplorasi 5", style="Title")
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED


def block_table(doc, stem, shift, truth, vals, crops_dir_e5, any_wrong_fn):
    t = doc.add_table(rows=2, cols=4)
    for b in range(8):
        cell = t.rows[b // 4].cells[b % 4]
        style_cell(cell, margin=(80, 80), valign="top")
        cell.width = Cm(6.7)
        p = cell_text(cell, f"Blok {b} · {JAM[shift][b]}", bold=True, size=9)
        if any_wrong_fn(b):
            r = p.add_run("   ada salah")
            r.font.size = Pt(8)
            r.font.color.rgb = RED
        img_path = crops_dir_e5 / f"{stem}_b{b}.jpg"
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        if img_path.exists():
            ih, iw = cv2.imread(str(img_path)).shape[:2]
            h_cm = min(1.15, 6.2 * ih / iw)  # muat di lebar sel (~6,2 cm)
            p2.add_run().add_picture(str(img_path), height=Cm(h_cm))
        else:
            r = p2.add_run("blok kosong — tidak dibaca model")
            r.font.size = Pt(8)
            r.font.color.rgb = MUTED
        nt = cell.add_table(rows=1 + 1 + len(ROWS), cols=7)
        hdr = ["", "1", "2", "3", "4", "5", "6"]
        for j, h in enumerate(hdr):
            c = nt.rows[0].cells[j]
            style_cell(c, fill=HEAD, margin=(10, 20))
            cell_text(c, h, size=7.5, align="center", mono=True)
        rows = [("Kunci", truth[b])] + [(lbl, vals[tag][b] if b < len(vals[tag]) else [None] * 6)
                                        for lbl, tag in ROWS]
        for i, (lbl, vv) in enumerate(rows, 1):
            c = nt.rows[i].cells[0]
            style_cell(c, fill=HEAD, margin=(10, 20))
            cell_text(c, lbl, size=7.5, bold=True, mono=True)
            for k in range(6):
                v = vv[k] if k < len(vv) else None
                g = truth[b][k]
                c = nt.rows[i].cells[k + 1]
                fill = None
                if lbl != "Kunci" and g != "SKIP":
                    fill = OK if v == g else BAD
                style_cell(c, fill=fill, margin=(10, 20))
                txt = "·" if v is None else ("(skip)" if v == "SKIP" else str(v))
                cell_text(c, txt, size=7.5, align="center", mono=True,
                          bold=(fill == BAD))
        for r in nt.rows:
            r.cells[0].width = Cm(1.1)
            for k in range(1, 7):
                r.cells[k].width = Cm(0.85)


def wrong_gallery(doc, stem, shift, truth, vals_e4, crops_dir_e4, conf=None, ambang=None):
    wrong = [(b, k) for b in range(8) for k in range(6)
             if truth[b][k] != "SKIP" and vals_e4[b][k] != truth[b][k]]
    doc.add_paragraph(f"Kotak yang salah di E4 ({len(wrong)}) — gambar = crop yang dibaca TrOCR E4",
                      style="Heading 2")
    if not wrong:
        p = doc.add_paragraph()
        r = p.add_run("Tidak ada — semua kotak benar.")
        r.font.color.rgb = MUTED
        return
    n_rows = (len(wrong) + 7) // 8
    t = doc.add_table(rows=n_rows, cols=8)
    for idx in range(n_rows * 8):
        cell = t.rows[idx // 8].cells[idx % 8]
        cell.width = Cm(3.3)
        if idx >= len(wrong):
            continue
        b, k = wrong[idx]
        style_cell(cell, valign="top")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img = crops_dir_e4 / f"{stem}_b{b}_k{k + 1}.jpg"
        if img.exists():
            p.add_run().add_picture(str(img), width=Cm(2.2))
        cell_text(cell, f"Blok {b} kotak {k + 1}", size=8, bold=True, align="center", first=False)
        g = truth[b][k]
        v = vals_e4[b][k]
        cell_text(cell, f"kunci: {'·' if g is None else g}", size=8, align="center", color=MUTED, first=False)
        pp = cell_text(cell, f"E4: {'·' if v is None else v}", size=8, align="center", color=RED,
                       mono=True, first=False)
        if norm(v) == norm(g):
            r = pp.add_run("  (huruf besar/kecil)")
            r.font.size = Pt(7)
            r.font.color.rgb = MUTED
        elif conf is not None and (stem, shift, b, k + 1) in conf:
            pst = conf[(stem, shift, b, k + 1)]
            ragu = pst < ambang
            cell_text(cell, (f"ditandai ragu · yakin {pst:.2f}" if ragu else f"LOLOS · yakin {pst:.2f}").replace(".", ","),
                      size=7, align="center", color=WARN if ragu else RED, bold=not ragu, first=False)


def overlay_for(photo, shift):
    d = np.load(HERE / "reference" / "ref_template.npz")
    anchors = json.loads((HERE / "reference" / "ref_anchors.json").read_text())
    det = detect_grid_fast(str(photo), shift, d["kp_pts"], d["des"], int(d["img_w"]), int(d["img_h"]), anchors)
    if not det.get("refined"):
        return None, det
    ov = draw_refined_overlay(det["corrected_image"], det["top_pts_px"], det["bot_pts_px"],
                              cell_ink=det["cell_ink"])
    s = 2400 / ov.shape[1]
    return cv2.resize(ov, None, fx=s, fy=s, interpolation=cv2.INTER_AREA), det


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--hasil", default=str(HERE / "output_bandingkan_e4_e5_ak"),
                    help="Hasil utama (mode kandidat, huruf a-k).")
    ap.add_argument("--hasil-lama", default=str(HERE / "output_bandingkan_e4_e5"),
                    help="Uji pertama (kandidat huruf a-z + mode bebas), utk baris pembanding di ringkasan.")
    ap.add_argument("--images-dir", default=str(HERE.parent.parent / "images"))
    ap.add_argument("--kunci", default=str(HERE.parent / "kunci_jawaban_eksplorasi.json"))
    ap.add_argument("--narasi", default=str(HERE / "narasi_laporan_e4_e5.json"),
                    help="Teks temuan/rekomendasi (JSON) -- ditulis setelah melihat hasil.")
    ap.add_argument("--out", default=str(HERE.parent.parent / "Laporan" / "Evaluasi Visual Eksplorasi 4 vs 5.docx"))
    args = ap.parse_args()

    hasil = Path(args.hasil)
    data = json.loads((hasil / "ringkasan.json").read_text(encoding="utf-8"))
    key = json.loads(Path(args.kunci).read_text(encoding="utf-8"))
    narasi = json.loads(Path(args.narasi).read_text(encoding="utf-8"))
    rows, block_rows, summ = data["per_foto"], data["per_blok"], data["ringkasan"]
    lama_path = Path(args.hasil_lama) / "ringkasan.json"
    if lama_path.exists():
        lama = json.loads(lama_path.read_text(encoding="utf-8"))["ringkasan"]
        for tag in ("e4_kandidat", "e5_kandidat"):
            if tag in lama:
                summ[tag + "_az"] = lama[tag]
        for tag in ("e4_bebas", "e5_bebas"):
            if tag in lama:
                summ[tag] = lama[tag]

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(1.5))
    doc.styles["Normal"].font.size = Pt(10)

    add_title(doc, "Ekstraksi tabel Lost time dari foto formulir — pembaca tulisan TrOCR (tanpa pelatihan) "
                   f"dibandingkan dengan Ollama · {date.today().strftime('%d-%m-%Y')}")

    doc.add_paragraph("Tentang dokumen ini", style="Heading 1")
    for b, t in narasi["tentang"]:
        bullet(doc, b, t)

    doc.add_paragraph("Ringkasan", style="Heading 1")
    p = doc.add_paragraph()
    r = p.add_run("Akurasi seluruh foto (12 foto × shift 1 & 2). \"Berisi\" = hanya kotak yang ada tulisannya; "
                  "kotak kosong diputuskan deteksi tinta yang sama di semua eksplorasi, jadi selalu benar. "
                  "\"Abaikan besar/kecil\" = cara aplikasi membandingkan kode (ProblemCodeConverter menormalkan ke huruf kecil).")
    r.font.size = Pt(9)
    r.font.color.rgb = MUTED
    srows = []
    for tag in ("e1", "e2", "e4_kandidat", "e5_kandidat", "e4_kandidat_az", "e5_kandidat_az",
                "e4_bebas", "e5_bebas"):
        if tag not in summ:
            continue
        s = summ[tag]
        srows.append([NAMA[tag], s["semua"], s["berisi"], s["berisi_ci"],
                      narasi["waktu"].get(tag, s["rata_waktu"])])
    simple_table(doc, ["Eksplorasi", "Semua kotak", "Kotak berisi", "Berisi (abaikan besar/kecil)",
                       "Waktu rata-rata / foto / shift"], srows, [9.0, 3.8, 3.8, 4.6, 5.0])
    for b_, t_ in narasi.get("ringkasan_catatan", []):
        bullet(doc, b_, t_)

    doc.add_paragraph("Per foto", style="Heading 2")
    prow = []
    for r_ in rows:
        stem, sh = r_["foto"], r_["shift"]
        def sc(tag):
            v = r_.get(f"{tag}_berisi_benar")
            f = frac(v)
            if not f or f[1] == 0:
                return "-"
            return (v, OK if f[0] == f[1] else (BAD if f[0] / f[1] < 0.75 else None))
        tun = " *" if stem in TUNING_FOTO else ""
        prow.append([stem + tun, sh, sheet_of(key, stem), r_.get("e4_kandidat_status", "-"),
                     sc("e1"), sc("e2"), sc("e4_kandidat"), f"{r_.get('e4_kandidat_total_dtk', '-')} dtk",
                     sc("e5_kandidat"), f"{r_.get('e5_kandidat_total_dtk', '-')} dtk"])
    simple_table(doc, ["Foto", "Shift", "Lembar", "Geometri", "E1 berisi", "E2 berisi", "E4 berisi", "E4 waktu",
                       "E5 berisi", "E5 waktu"], prow, [4.6, 1.2, 1.6, 3.2, 2.2, 2.2, 2.2, 2.2, 2.2, 2.2])
    p = doc.add_paragraph()
    r = p.add_run("* foto yang dipakai saat menyetel praproses (lihat bagian Pengembangan). "
                  "Angka = kotak berisi yang dibaca persis sama dengan kunci.")
    r.font.size = Pt(8)
    r.font.color.rgb = MUTED

    for sec_name in ("pengembangan", "temuan", "hibrida", "rekomendasi"):
        if sec_name not in narasi:
            continue
        doc.add_paragraph(narasi[sec_name]["judul"], style="Heading 1")
        for item in narasi[sec_name]["isi"]:
            if isinstance(item, dict) and "subjudul" in item:
                doc.add_paragraph(item["subjudul"], style="Heading 2")
            elif isinstance(item, dict) and "gambar" in item:
                imgs = [hasil / g for g in item["gambar"]]
                t_ = doc.add_table(rows=1, cols=len(imgs))
                for j, im in enumerate(imgs):
                    c_ = t_.rows[0].cells[j]
                    c_.width = Cm(max(item.get("lebar_cm", 2.5) + 0.6, 4.2))
                    pp_ = c_.paragraphs[0]
                    pp_.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if im.exists():
                        pp_.add_run().add_picture(str(im), width=Cm(item.get("lebar_cm", 2.5)))
                    if item.get("label"):
                        cell_text(c_, item["label"][j], size=8, align="center", color=MUTED, first=False)
                if item.get("keterangan"):
                    p_ = doc.add_paragraph()
                    r_k = p_.add_run(item["keterangan"])
                    r_k.font.size = Pt(8)
                    r_k.font.color.rgb = MUTED
            elif isinstance(item, dict) and "tabel" in item:
                simple_table(doc, item["tabel"][0], item["tabel"][1:], item["lebar"])
                doc.add_paragraph()
            elif isinstance(item, list):
                bullet(doc, item[0], item[1])
            else:
                doc.add_paragraph(item)

    conf, ambang = None, narasi.get("ambang_hibrida")
    hib = hasil / "hibrida" / "hibrida_e4.json"
    if hib.exists() and ambang is not None:
        conf = {(c["foto"], c["shift"], c["blok"], c["kolom"]): c["posterior"]
                for c in json.loads(hib.read_text(encoding="utf-8"))["sel"]}

    photos = {Path(p).name.split(".")[0]: p for p in Path(args.images_dir).glob("*.jp*g")}
    for r_ in rows:
        stem, sh = r_["foto"], r_["shift"]
        truth = truth_of(block_rows, stem, sh)
        if not truth or all(v is None for blk in truth for v in blk):
            continue  # shift kosong semua -- spt laporan lama, tdk ditampilkan
        vals = {tag: block_vals(block_rows, stem, sh, tag) for _, tag in ROWS}
        doc.add_page_break()
        doc.add_paragraph(f"{stem} — shift {sh}", style="Heading 1")
        ov, det = overlay_for(photos[stem], sh)
        geo = f"{det.get('status')} ({det.get('refine', {}).get('quality')})"
        p = doc.add_paragraph()
        parts = [f"Lembar {sheet_of(key, stem)}   ·   geometri: {geo}"]
        for lbl, tag in ROWS:
            v = r_.get(f"{tag}_berisi_benar")
            t = r_.get(f"{tag}_total_dtk")
            parts.append(f"{lbl} {v or '-'}" + (f" ({t} dtk)" if t is not None else ""))
        r = p.add_run("   ·   ".join(parts) + ("   ·   foto tuning" if stem in TUNING_FOTO else ""))
        r.font.size = Pt(10)
        r.font.color.rgb = WARN if det.get("status") != "success" else MUTED
        if ov is not None:
            doc.add_paragraph().add_run().add_picture(img_stream(ov), width=Cm(26.5))

        doc.add_paragraph("Per blok jam (gambar = baris yang dibaca TrOCR E5; kotak kosong sudah dibuang)",
                          style="Heading 2")

        def any_wrong(b):
            return any(truth[b][k] != "SKIP" and vals[tag][b][k] != truth[b][k]
                       for _, tag in ROWS[2:] for k in range(6))
        block_table(doc, stem, sh, truth, vals,
                    hasil / "crops" / "e5_kandidat" / f"{stem}_s{sh}", any_wrong)
        wrong_gallery(doc, stem, sh, truth, vals["e4_kandidat"],
                      hasil / "crops" / "e4_kandidat" / f"{stem}_s{sh}", conf, ambang)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Laporan -> {out.resolve()}")


if __name__ == "__main__":
    main()
