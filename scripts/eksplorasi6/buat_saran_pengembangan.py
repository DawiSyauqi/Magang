#!/usr/bin/env python3
"""
buat_saran_pengembangan.py -- docx TAMBAHAN SARAN PENGEMBANGAN (lanjutan Subbab
Saran pada laporan KP) berdasarkan eksperimen Eksplorasi 1-6. Laporan KP asli
TIDAK disentuh; format mengikuti laporan (Times New Roman 12, spasi 1,5, rata
kiri-kanan, istilah asing miring, butir bernomor dgn kalimat pembuka tebal).

Penanda teks: **tebal**, *miring*.

    ../eksplorasi4/.venv/Scripts/python buat_saran_pengembangan.py
"""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).parent
NOMOR_AWAL = 3  # laporan KP sudah memuat 2 butir saran

PEMBUKA = (
    "Selain saran di atas, serangkaian eksperimen lanjutan terhadap modul AI Baca Kertas yang dilakukan "
    "setelah implementasi menghasilkan beberapa saran pengembangan berikut. Eksperimen dilakukan pada "
    "12 foto area tabel dan 7 foto seluruh halaman dari 2 lembar formulir yang isinya telah dicatat sebagai "
    "kunci jawaban (*ground truth*), dan dijalankan pada CPU laptop tanpa GPU. Karena jumlah lembar masih "
    "sedikit, angka-angka di bawah ini merupakan indikasi awal yang perlu dikonfirmasi pada data yang lebih banyak."
)

SARAN = [
    ("Penggunaan model pengenal tulisan tangan khusus untuk kotak waktu tidak produktif.",
     "Isi setiap kotak 10 menit hanya berupa kode pendek (angka 0–8 dengan huruf a–k, atau penanda “x”), "
     "sehingga kotak dapat dibaca dengan model pengenal tulisan tangan berukuran kecil, yaitu TrOCR "
     "(*Transformer-based Optical Character Recognition*) varian *small-handwritten*, yang diminta memilih kode "
     "paling sesuai dari daftar kode yang sah, bukan menulis jawaban bebas. Pada eksperimen, cara ini membaca "
     "benar sekitar 90–91% kotak terisi, dibandingkan sekitar 54% pada pembacaan per kotak dan sekitar 86% pada "
     "pembacaan per blok jam menggunakan Qwen2.5-VL melalui Ollama. Model ini dapat dijalankan pada CPU dengan "
     "waktu sekitar 0,8 detik per kotak terisi tanpa pelatihan ulang, sehingga tidak bergantung pada GPU server."),
    ("Pelurusan tabel secara otomatis dengan pencocokan titik fitur cetakan formulir.",
     "Sebagai alternatif atau pelengkap model deteksi objek seperti YOLO-OBB, posisi tabel dapat ditemukan "
     "dengan mencocokkan titik-titik fitur khas pada bagian cetakan formulir (judul, label, garis, dan legenda) "
     "terhadap satu foto acuan menggunakan SIFT (*Scale-Invariant Feature Transform*) dan homografi dengan "
     "penyaringan RANSAC, lalu menyesuaikan posisi baris ke garis cetak tabel yang sebenarnya. Cara ini tidak "
     "memerlukan data pelatihan dan pada eksperimen berhasil memetakan seluruh foto uji (12 dari 12 foto area "
     "tabel dan 7 dari 7 foto seluruh halaman), termasuk foto yang diputar 90° dan foto yang diambil dari jarak "
     "jauh, sehingga penandaan tiga titik sudut oleh petugas dapat dihilangkan. Peta posisi formulir cukup "
     "dibuat satu kali dan perlu dibuat ulang hanya apabila tata letak formulir berubah."),
    ("Penyederhanaan alur menjadi satu foto seluruh halaman.",
     "Eksperimen pembacaan seluruh halaman dari satu foto menunjukkan bahwa akurasi kotak waktu tidak produktif "
     "tidak menurun dibandingkan foto khusus area tabel (sekitar 90% berbanding 91%). Pada foto yang sama, "
     "tanggal, kode mesin, shift, ukuran (*size*), dan kecepatan (*speed*) terbaca benar pada 7 dari 7 foto, "
     "dengan waktu proses sekitar 41–77 detik per foto pada CPU laptop. Disarankan Section 1 dan Section 2 "
     "digabung menjadi satu kali pemotretan, dengan hasil bacaan header ditampilkan sebagai isian awal yang tetap "
     "dapat dikoreksi petugas. Dengan rancangan ini, apabila pembacaan header ternyata kurang andal pada lembar "
     "lain, header cukup diisi manual tanpa mengubah bagian pembacaan tabel."),
    ("Penandaan kotak yang meragukan beserta kandidat jawabannya.",
     "Setiap kotak yang dibaca dapat diberi nilai keyakinan. Kotak dengan keyakinan di bawah ambang tertentu "
     "(pada eksperimen 0,85) ditandai pada layar tinjauan bersama tiga kandidat kode teratas, sehingga petugas "
     "cukup memilih salah satunya. Pada eksperimen, sekitar 20% kotak terisi ditandai dan sebagian besar "
     "kesalahan tertangkap, sehingga hasil akhir setelah tinjauan mencapai sekitar 98%. Cara ini membuat "
     "peninjauan petugas (*human-in-the-loop*) lebih terarah karena perhatian difokuskan pada kotak yang "
     "benar-benar meragukan."),
    ("Validasi otomatis pada field header dan pengisian nama operator dari data master.",
     "Tanggal hasil bacaan dapat diperiksa keabsahannya serta kecocokan nama harinya dengan kalender; pada "
     "eksperimen, pemeriksaan ini menemukan formulir bertuliskan “Senin, 3-08-2025” padahal tanggal tersebut "
     "jatuh pada hari Minggu. Sebaliknya, nama operator yang ditulis tangan, terlebih yang disertai coretan, "
     "belum dapat dibaca dengan andal (terbaca menjadi kata lain). Oleh karena itu, nama operator dan kode mesin "
     "disarankan dipilih dari daftar master (TMMESINOP dan MFRESMAS) atau dari akun petugas yang masuk, dengan "
     "hasil bacaan model hanya digunakan sebagai pilihan awal."),
    ("Layanan Python yang tetap berjalan untuk mempersingkat waktu proses.",
     "Pada eksperimen, setiap pemanggilan skrip sebagai proses terpisah menghabiskan sekitar 15 detik hanya untuk "
     "menyalakan Python, memuat pustaka, dan memuat model, di luar waktu membaca kotak. Disarankan pembaca "
     "dijalankan sebagai layanan lokal yang memuat model satu kali saat server dinyalakan dan dipanggil oleh "
     "Laravel melalui permintaan HTTP intranet, sehingga waktu tunggu per foto dapat berkurang tanpa mengubah "
     "akurasi."),
    ("Pengujian ulang pada lembar formulir baru sebelum diterapkan di produksi.",
     "Seluruh eksperimen di atas baru menggunakan 2 lembar formulir dengan penulis yang kemungkinan sama, dan "
     "sebagian aturan pembacaan header disusun sambil melihat foto yang sama. Sebelum diterapkan, pendekatan ini "
     "perlu diuji pada lembar dan penulis yang berbeda tanpa mengubah pengaturan, termasuk kondisi yang belum "
     "diuji seperti pencahayaan kurang, foto buram, dan formulir yang terpotong di tepi foto. Pengujian ini "
     "dapat sekaligus menjadi bagian dari penyusunan *dataset* beranotasi pada saran pertama."),
]


def runs(p, text, bold=False):
    """Tulis teks dgn penanda **tebal** / *miring*."""
    for part in re.split(r"(\*\*.+?\*\*|\*.+?\*)", text):
        if not part:
            continue
        if part.startswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("*"):
            r = p.add_run(part[1:-1])
            r.italic = True
            r.bold = bold
        else:
            r = p.add_run(part)
            r.bold = bold


def fmt(p, first_line=None, left=None, hanging=None):
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if left is not None:
        pf.left_indent = left
    if hanging is not None:
        pf.first_line_indent = -hanging
    elif first_line is not None:
        pf.first_line_indent = first_line


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE.parent.parent / "Laporan" / "Saran Pengembangan - Hasil Eksplorasi.docx"))
    args = ap.parse_args()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Cm(2.54)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(12)
    st.font.color.rgb = RGBColor(0, 0, 0)

    h = doc.add_paragraph()
    fmt(h)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = h.add_run("Saran Pengembangan Berdasarkan Eksperimen Lanjutan")
    r.bold = True

    p = doc.add_paragraph()
    fmt(p, first_line=Cm(0.64))
    runs(p, PEMBUKA)

    for i, (judul, isi) in enumerate(SARAN, start=NOMOR_AWAL):
        p = doc.add_paragraph()
        fmt(p, left=Cm(1.26), hanging=Cm(0.63))
        p.paragraph_format.tab_stops.add_tab_stop(Cm(1.26))
        p.add_run(f"{i}.\t")
        runs(p, judul, bold=True)
        p.add_run(" ")
        runs(p, isi)

    out = Path(args.out)
    doc.save(out)
    print(f"Saran -> {out.resolve()}")


if __name__ == "__main__":
    main()
