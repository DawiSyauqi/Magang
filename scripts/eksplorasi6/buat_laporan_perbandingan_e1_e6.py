#!/usr/bin/env python3
"""
buat_laporan_perbandingan_e1_e6.py -- laporan PERBANDINGAN & REKOMENDASI
Eksplorasi 1 s.d. 6 (docx, gaya sama dgn laporan evaluasi lain).

Angka diambil dari hasil yg sudah ada (tidak menjalankan model):
  E1/E2      : ../output_bandingkan_eksplorasi (via ringkasan E4/E5 _ak)
  E4/E5      : ../eksplorasi4/output_bandingkan_e4_e5_ak (+ hibrida/)
  E6         : output_e6/ringkasan.json
  E3         : tidak ada hasil terukur (lihat narasi)
Teks: narasi_perbandingan_e1_e6.json.

    ../eksplorasi4/.venv/Scripts/python buat_laporan_perbandingan_e1_e6.py
"""

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from buat_laporan_docx_e6 import BAD, HEAD, INK, MUTED, OK, bullet, cell_text, para, simple_table, style_cell  # noqa: E402

WARNF = "FFF4D6"


def fill_of(tag):
    return {"ok": OK, "bad": BAD, "warn": WARNF, "head": HEAD}.get(tag)


def table_from(doc, spec):
    rows = [[(c[0], fill_of(c[1])) if isinstance(c, list) else c for c in r] for r in spec["baris"]]
    simple_table(doc, spec["kolom"], rows, spec["lebar_cm"])
    if spec.get("catatan"):
        para(doc, spec["catatan"], color=MUTED, size=8.5, italic=True)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--narasi", default=str(HERE / "narasi_perbandingan_e1_e6.json"))
    ap.add_argument("--out", default=str(HERE.parent.parent / "Laporan" / "Perbandingan dan Rekomendasi Eksplorasi 1-6.docx"))
    args = ap.parse_args()
    nar = json.loads(Path(args.narasi).read_text(encoding="utf-8"))

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.font.color.rgb = INK

    doc.add_paragraph(nar["judul"], style="Title")
    para(doc, nar["subjudul"], color=MUTED, size=11)

    for sec_ in nar["bagian"]:
        if sec_.get("halaman_baru"):
            doc.add_page_break()
        doc.add_paragraph(sec_["judul"], style="Heading 1")
        for item in sec_["isi"]:
            kind = item["jenis"]
            if kind == "paragraf":
                para(doc, item["teks"])
            elif kind == "sub":
                doc.add_paragraph(item["teks"], style="Heading 2")
            elif kind == "poin":
                for b in item["daftar"]:
                    bullet(doc, b[0], b[1])
            elif kind == "tabel":
                table_from(doc, item)
            elif kind == "kotak":  # kotak kesimpulan berwarna
                t = doc.add_table(rows=1, cols=1)
                c = t.rows[0].cells[0]
                style_cell(c, fill=fill_of(item.get("warna", "ok")), margin=(120, 160))
                cell_text(c, item["judul"], bold=True, size=11)
                for line in item["baris"]:
                    cell_text(c, line, size=10, first=False)
                c.width = Cm(26.5)
                para(doc, "")

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"Laporan -> {out.resolve()}")


if __name__ == "__main__":
    main()
