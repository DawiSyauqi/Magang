#!/usr/bin/env python3
"""
buat_laporan_alur_e1_e6.py -- docx "Alur Kerja Eksplorasi 1-6 & Peran OCR
Tesseract": alur tiap eksplorasi tahap demi tahap + di mana Tesseract dipakai
(OSD orientasi & OCR-anchor label "Jam"), disusun dari isi skrip:
  eksplorasi1 : baca_kotak_eksplorasi.py, geometry_pipeline.py, detect_grid.py,
                reference_calib.py, ocr_anchor.py, imgutil.py
  eksplorasi2 : ollama_block_pipeline_refine_eksplorasi.py, crop_blocks*.py
  eksplorasi3 : ollama_full_photo_pipeline.py, preprocess.py
  eksplorasi4/5 : baca_kotak_trocr_eksplorasi.py / baca_blok_trocr_eksplorasi.py,
                refine_grid_eksplorasi.py (detect_grid_fast + fallback)
  eksplorasi6 : baca_fullpage_eksplorasi.py, geometri_fullpage.py
Angka hasil dari laporan yg sudah ada (tidak menjalankan model).

    ../eksplorasi4/.venv/Scripts/python buat_laporan_alur_e1_e6.py
"""

import argparse
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Cm, Pt

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
sys.dont_write_bytecode = True
from buat_laporan_docx_e6 import BAD, HEAD, INK, MUTED, OK, bullet, cell_text, para, style_cell  # noqa: E402

WARNF = "FFF4D6"
FILL = {"ok": OK, "bad": BAD, "warn": WARNF, None: None}


def table(doc, header, rows, widths):
    """Tabel rata kiri; sel boleh (teks, 'ok'|'bad'|'warn')."""
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        style_cell(c, fill=HEAD)
        cell_text(c, h, bold=True, size=9)
    for i, row in enumerate(rows, 1):
        for j, v in enumerate(row):
            fill = None
            if isinstance(v, tuple):
                v, fill = v
            c = t.rows[i].cells[j]
            style_cell(c, fill=FILL[fill], valign="top")
            cell_text(c, v, size=9, bold=(j == 0))
    for j, w in enumerate(widths):
        for r in t.rows:
            r.cells[j].width = Cm(w)
    para(doc, "")
    return t


def kotak(doc, judul, baris, warna="ok"):
    t = doc.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    style_cell(c, fill=FILL[warna], margin=(120, 160))
    cell_text(c, judul, bold=True, size=11)
    for line in baris:
        cell_text(c, line, size=10, first=False)
    c.width = Cm(26.5)
    para(doc, "")


TESS_YA, TESS_CADANGAN, TESS_TIDAK = ("Ya, tiap foto", "bad"), ("Hanya cadangan", "warn"), ("Tidak", "ok")

EKSPLORASI = [
    {
        "judul": "Eksplorasi 1 — close-up grid, Ollama membaca per kotak",
        "ide": "Mengotomatiskan pencarian posisi grid (menggantikan klik 3 titik), lalu membaca tiap kotak 10 menit "
               "satu per satu dengan Ollama (qwen2.5vl:3b), mirip Section 2 produksi.",
        "input": "Foto close-up grid Lost time + shift (1/2/3).",
        "tahap": [
            ("1. Posisi grid (jalur cepat)", "Titik fitur ORB foto dicocokkan ke foto referensi 134505 → homography. "
             "Rotasi foto ditebak dari homography (tanpa OSD). Letak kasar baris shift diambil dari referensi.",
             "OpenCV (ORB, RANSAC)"),
            ("2. Rapikan ke garis cetak", "refine_grid mencari garis cetak tabel yang sebenarnya di sekitar letak kasar, "
             "jadi kemiringan & kertas melengkung ikut tertangani.", "OpenCV"),
            ("   Jalur cadangan", "Bila langkah 1–2 gagal: foto diluruskan dengan OSD Tesseract, lalu posisi baris "
             "dicek dengan OCR label cetak \"Jam\" (detect_grid lama) + refine.", "Tesseract (OSD + OCR \"Jam\")"),
            ("3. Cek tinta", "48 kotak diukur proporsi tintanya; kotak kosong langsung null tanpa memanggil model.", "OpenCV"),
            ("4. Potong per kotak", "Tiap kotak berisi di-warp perspektif dari 4 sudutnya (lurus).", "OpenCV"),
            ("5. Baca", "1 kotak = 1 panggilan Ollama.", "Ollama qwen2.5vl:3b"),
            ("6. Susun", "48 hasil disusun menjadi 8 blok jam × 6 kotak.", "Python"),
        ],
        "tesseract": "Versi akhir (yang dinilai): hanya di jalur cadangan. Versi awal E1 (geometry_pipeline.py, "
                     "murni geometri tanpa membaca isi) memakai Tesseract di SETIAP foto: OSD untuk meluruskan foto, "
                     "lalu OCR \"Jam\" untuk mengecek posisi baris hasil homography.",
        "hasil": "54,2% kotak berisi benar; geometri 12/12; ±33 dtk (lembar jarang) – ±155 dtk (lembar padat) per shift.",
    },
    {
        "judul": "Eksplorasi 2 — close-up grid, Ollama membaca per blok jam",
        "ide": "Geometri sama dengan E1, tetapi Ollama membaca 6 kotak sekaligus (1 blok jam) supaya panggilan model "
               "lebih sedikit.",
        "input": "Foto close-up grid Lost time + shift.",
        "tahap": [
            ("1–2. Posisi grid", "Sama dengan E1: homography ORB ke referensi + refine_grid; jalur cadangan OSD + OCR "
             "\"Jam\".", "OpenCV; Tesseract (cadangan)"),
            ("3. Cek tinta (ink gate)", "Blok yang 6 kotaknya kosong tidak dikirim ke model; jawaban model untuk kotak "
             "kosong dipaksa null.", "OpenCV"),
            ("4. Potong per blok", "Blok di-warp lurus dengan garis bantu; prompt menyebut kolom mana yang bertinta "
             "(ink hint).", "OpenCV"),
            ("5. Baca", "1 blok = 1 panggilan Ollama, jawaban JSON 6 kotak.", "Ollama qwen2.5vl:3b"),
        ],
        "tesseract": "Hanya di jalur cadangan (sama dengan E1). Versi E2 paling awal (ollama_block_pipeline.py + "
                     "crop_blocks.py) memakai detect_grid lama: OSD + OCR \"Jam\" di setiap foto.",
        "hasil": "85,7% kotak berisi benar; ±42–90 dtk per shift.",
    },
    {
        "judul": "Eksplorasi 3 — foto utuh tanpa crop, Ollama mencari sendiri",
        "ide": "Tidak ada geometri sama sekali: model diberi foto utuh dan diminta menemukan baris shift & blok jam "
               "sendiri.",
        "input": "Foto (utuh) + shift.",
        "tahap": [
            ("1. Luruskan orientasi", "OSD Tesseract menebak rotasi foto (0/90/180/270°) lalu foto diputar. "
             "Bila OSD gagal → needs_retake.", "Tesseract (OSD)"),
            ("2. Praproses", "Upscale bila kecil + enhance kontras (fungsi dari paper_reader_extract.py).", "OpenCV"),
            ("3. Baca", "Tiap panggilan Ollama menerima foto utuh yang sama; model harus menemukan blok yang dimaksud.",
             "Ollama"),
        ],
        "tesseract": "Ya, di SETIAP foto (OSD untuk meluruskan orientasi). Tanpa Tesseract, E3 tidak bisa jalan.",
        "hasil": "Tidak ada hasil terukur dengan kunci jawaban; laporan KP mencatat foto utuh gagal untuk kotak kecil.",
    },
    {
        "judul": "Eksplorasi 4 — close-up grid, TrOCR membaca per kotak",
        "ide": "Tiruan E1, hanya pembacanya yang diganti: TrOCR small-handwritten (tanpa pelatihan, CPU, tanpa Ollama) "
               "yang memilih kode paling cocok dari daftar kode sah.",
        "input": "Foto close-up grid Lost time + shift.",
        "tahap": [
            ("1–2. Posisi grid", "Sama dengan E1 (salinan refine_grid_eksplorasi.py): homography ORB + refine_grid; "
             "jalur cadangan OSD + OCR \"Jam\".", "OpenCV; Tesseract (cadangan)"),
            ("3. Cek tinta", "Kotak kosong = null tanpa model.", "OpenCV"),
            ("4. Potong per kotak", "Warp per kotak, pita tepi 8% (sisa garis kotak) diisi warna kertas, kontras "
             "diregangkan.", "OpenCV"),
            ("5. Baca", "Semua kode sah (0–8 + huruf a–k, atau x) dinilai TrOCR, ambil skor tertinggi.",
             "TrOCR small"),
            ("6. Konfirmasi (hibrida)", "Keyakinan < 0,85 → kotak ditandai ragu + 3 kandidat.", "Python"),
        ],
        "tesseract": "Hanya di jalur cadangan. Dalam pengujian, 24 dari 24 proses (12 foto × 2 shift) lewat jalur cepat, "
                     "jadi Tesseract tidak pernah terpanggil.",
        "hasil": "91,0% kotak berisi benar; dengan konfirmasi 98,7%; ±22–57 dtk per shift.",
    },
    {
        "judul": "Eksplorasi 5 — close-up grid, TrOCR membaca per blok jam",
        "ide": "Tiruan E2 dengan TrOCR: 1 blok dibaca sebagai 1 baris teks.",
        "input": "Foto close-up grid Lost time + shift.",
        "tahap": [
            ("1–2. Posisi grid", "Sama dengan E2/E4; jalur cadangan OSD + OCR \"Jam\".", "OpenCV; Tesseract (cadangan)"),
            ("3. Cek tinta", "Blok kosong tidak dibaca; kotak kosong dibuang dari gambar, kotak bertinta dirangkai "
             "rapat.", "OpenCV"),
            ("4. Potong per blok", "Label jam dan garis cetak dibuang (TrOCR akan ikut membacanya).", "OpenCV"),
            ("5. Baca", "Beam search dibatasi pola \"<kode> <kode> …\" tepat sejumlah kotak bertinta.", "TrOCR small"),
        ],
        "tesseract": "Hanya di jalur cadangan (sama dengan E4).",
        "hasil": "74,4% kotak berisi benar; ±22–41 dtk per shift.",
    },
    {
        "judul": "Eksplorasi 6 — foto seluruh halaman, header + Lost time sekaligus",
        "ide": "Satu foto utuh menggantikan Section 1 + Section 2: header, Size/Speed, operator dan Lost time dibaca "
               "sekaligus, tanpa klik titik.",
        "input": "Foto seluruh halaman (shift boleh otomatis dari header).",
        "tahap": [
            ("1. Daftarkan ke formulir", "Titik fitur SIFT foto dicocokkan ke kanvas referensi (hanya bagian CETAKAN "
             "formulir; area isian dimasker) → homography RANSAC → foto diluruskan jadi kanvas 3400 × 2330 px. "
             "SIFT tahan rotasi & skala, jadi foto miring/diputar 90° langsung terpetakan.", "OpenCV (SIFT)"),
            ("2. Header & Size/Speed", "Kotak field di kanvas dipotong; garis cetak & coretan tebal dibuang; field "
             "kosong dilewati. Tanggal & mesin dibaca TrOCR dengan pola (regex); angka Size/Speed dibaca per "
             "karakter; shift dari 3 kandidat; tanggal dicek ke kalender.", "OpenCV + TrOCR"),
            ("3. Lost time", "Baris tiap shift di-snap ke garis cetak (refine_grid), cek tinta, kotak dibaca seperti "
             "E4, keyakinan < 0,85 ditandai ragu.", "OpenCV + TrOCR"),
        ],
        "tesseract": "Tidak dipakai. Referensi E6 dibuat dari bingkai formulir & titik fitur cetakan, bukan OCR. "
                     "(Modul imgutil.py tetap mengimpor paket pytesseract, tetapi fungsinya tidak dipanggil.)",
        "hasil": "Header 7/7 (kecuali operator 3/7); Lost time 90,0%, dengan konfirmasi 98,2%; ±41–77 dtk per foto.",
    },
]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=str(HERE.parent.parent / "Laporan" / "Alur Eksplorasi 1-6 dan Peran Tesseract.docx"))
    args = ap.parse_args()

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.6)
    sec.top_margin = sec.bottom_margin = Cm(1.4)
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10)
    st.font.color.rgb = INK

    doc.add_paragraph("Alur Kerja Eksplorasi 1–6 dan Peran OCR Tesseract", style="Title")
    para(doc, "Disusun dari isi skrip di scripts/eksplorasi1 – eksplorasi6 · angka hasil dari laporan evaluasi yang "
              "sudah ada · 12 Oktober 2026", color=MUTED, size=11)

    # ------------------------------------------------------------------ jawaban singkat
    doc.add_paragraph("Jawaban singkat: Tesseract dipakai untuk apa?", style="Heading 1")
    kotak(doc, "Tesseract TIDAK pernah membaca tulisan tangan. Tugasnya hanya 2, keduanya untuk GEOMETRI:", [
        "1. OSD (Orientation & Script Detection) — menebak foto terbalik/miring 90°, lalu foto diputar tegak.",
        "2. OCR-anchor — mencari tulisan CETAK \"Jam\" (dan \"Lost time\") di kolom label kiri formulir untuk "
        "menentukan posisi baris shift 1/2/3.",
        "Isi kotak Lost time & header selalu dibaca model lain: Ollama (qwen2.5vl) di E1–E3 & produksi, TrOCR di E4–E6.",
    ])
    para(doc, "Kapan Tesseract dipakai di tiap eksplorasi:")
    table(doc, ["Eksplorasi / proses", "Pakai Tesseract?", "Untuk apa"], [
        ["Kalibrasi referensi E1–E5 (sekali, foto 134505)", ("Ya, sekali saja", "warn"),
         "OSD + OCR \"Jam\" → posisi baris shift disimpan ke reference/ref_anchors.json. File ini sudah ada, "
         "jadi tidak perlu diulang kecuali referensi dibuat ulang."],
        ["E1 versi awal (geometry_pipeline.py)", TESS_YA, "OSD meluruskan foto + OCR \"Jam\" mengecek posisi baris."],
        ["E1, E2 versi akhir · E4 · E5", TESS_CADANGAN,
         "Jalur cepat (ORB + refine) tanpa Tesseract. Tesseract dipanggil hanya kalau jalur cepat gagal "
         "(path = fallback_osd_ocr)."],
        ["E3", TESS_YA, "OSD meluruskan foto sebelum dikirim ke Ollama; OSD gagal → minta foto ulang."],
        ["E6", TESS_TIDAK, "Posisi dari SIFT + bingkai formulir."],
        ["Alur produksi (Section 1 & 2)", TESS_TIDAK,
         "Section 2 memakai 3 titik yang diklik petugas; tidak ditemukan pemanggilan Tesseract di "
         "paper_reader_extract.py."],
    ], [7.5, 4.0, 15.0])

    # ------------------------------------------------------------------ istilah
    doc.add_page_break()
    doc.add_paragraph("Istilah yang dipakai", style="Heading 1")
    for b, t in [
        ("Foto referensi:", "satu foto contoh formulir yang posisi grid/field-nya sudah diketahui; foto baru "
                            "dipetakan ke foto ini. E1–E5: foto close-up 134505; E6: foto utuh 143018."),
        ("ORB / SIFT + homography:", "mencari titik-titik khas yang sama di foto baru & referensi, lalu menghitung "
                                     "transformasi (geser, putar, skala, perspektif) di antara keduanya. Semuanya OpenCV."),
        ("refine_grid:", "setelah letak kasar diketahui, garis tabel dicari ulang di foto sebenarnya supaya "
                         "potongan kotak tepat di garis cetak."),
        ("Cek tinta (ink gate):", "mengukur proporsi piksel gelap tiap kotak; kotak kosong tidak dikirim ke model."),
        ("Ollama (qwen2.5vl):", "model bahasa-visual besar, dijalankan lokal; dipakai produksi, E1, E2, E3."),
        ("TrOCR small-handwritten:", "model pengenal tulisan tangan kecil dari Microsoft, jalan di CPU; dipakai E4–E6."),
    ]:
        bullet(doc, b, t)

    doc.add_paragraph("Ringkasan alur", style="Heading 1")
    table(doc, ["", "Foto", "Cara menemukan posisi kotak", "Pembaca isi", "Satuan baca", "Tesseract"], [
        ["E1", "close-up grid", "ORB + refine_grid", "Ollama", "per kotak", TESS_CADANGAN],
        ["E2", "close-up grid", "ORB + refine_grid", "Ollama", "per blok jam", TESS_CADANGAN],
        ["E3", "utuh", "tidak ada (model mencari sendiri)", "Ollama", "per blok jam", TESS_YA],
        ["E4", "close-up grid", "ORB + refine_grid", "TrOCR", "per kotak", TESS_CADANGAN],
        ["E5", "close-up grid", "ORB + refine_grid", "TrOCR", "per blok jam", TESS_CADANGAN],
        ["E6", "seluruh halaman", "SIFT ke kanvas formulir + refine_grid", "TrOCR", "per kotak + field header", TESS_TIDAK],
        ["Produksi", "utuh + close-up", "fraksi tetap + 3 titik diklik petugas", "Ollama", "per kotak", TESS_TIDAK],
    ], [2.0, 3.5, 8.0, 3.0, 5.0, 5.0])

    # ------------------------------------------------------------------ per eksplorasi
    for i, e in enumerate(EKSPLORASI):
        doc.add_page_break() if i % 2 == 0 else None
        doc.add_paragraph(e["judul"], style="Heading 1")
        bullet(doc, "Ide:", e["ide"])
        bullet(doc, "Masukan:", e["input"])
        table(doc, ["Tahap", "Proses", "Alat"], [list(t) for t in e["tahap"]], [4.5, 16.5, 5.5])
        bullet(doc, "Peran Tesseract:", e["tesseract"])
        bullet(doc, "Hasil (dari laporan evaluasi):", e["hasil"])

    # ------------------------------------------------------------------ kapan install
    doc.add_page_break()
    doc.add_paragraph("Jadi, perlu memasang Tesseract atau tidak?", style="Heading 1")
    table(doc, ["Kebutuhan", "Perlu Tesseract?", "Keterangan"], [
        ["Menjalankan E6 (/test/eksplorasi)", TESS_TIDAK, "—"],
        ["Menjalankan E4 / E5 dengan foto normal", TESS_TIDAK,
         "Jalur cepat sudah cukup (24/24 di pengujian; uji halaman /test/eksplorasi juga berhasil tanpa Tesseract)."],
        ["E1/E2/E4/E5 untuk foto yang gagal jalur cepat", ("Ya", "warn"),
         "Tanpa Tesseract, jalur cadangan error dan foto berakhir gagal."],
        ["Menjalankan E3", ("Ya", "bad"), "OSD dipanggil di setiap foto."],
        ["Membuat ulang referensi E1–E5 (calibrate_reference_cli.py)", ("Ya", "bad"), "OSD + OCR \"Jam\"."],
        ["Membuat referensi E6 (geometri_fullpage.py --buat-referensi)", TESS_TIDAK, "—"],
        ["Alur produksi sekarang", TESS_TIDAK, "—"],
    ], [8.5, 4.0, 14.0])
    para(doc, "Catatan: paket Python pytesseract tetap harus terpasang di venv karena diimpor oleh imgutil.py "
              "(sudah ada di scripts/eksplorasi4/.venv). Yang belum terpasang di laptop ini adalah program "
              "Tesseract-nya sendiri (tesseract.exe). Bila perlu: pasang installer Windows dari UB Mannheim ke "
              "C:\\Program Files\\Tesseract-OCR — skrip otomatis mencarinya di sana.", color=MUTED, size=9, italic=True)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f"Laporan -> {out.resolve()}")


if __name__ == "__main__":
    main()
